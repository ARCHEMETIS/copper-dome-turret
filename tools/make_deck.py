# =============================================================
# make_deck.py — สร้างเดคนำเสนอ 3 นาที (.pptx) จากเนื้อหาใน docs/slide-copy.md
# รัน:  venv\Scripts\python.exe tools\make_deck.py
#
# ทำไมเป็น pptx ไม่ใช่ Canva: Canva Magic Design ตอบ design_generation_error
# (ตัวกรองเนื้อหาคำว่า turret/ballistics หรือโควตา AI หมด) — pptx เปิดใน
# PowerPoint แก้ได้ทันที และลาก import เข้า Canva ได้ถ้าอยากได้ธีมสวยกว่านี้
#
# ตัวเลขทุกตัวในไฟล์นี้ตรงกับ src/config.py + docs/slide-assets/ballistics-output.txt
# ถ้าแก้ค่าในโค้ดจริง อย่าลืมแก้ที่นี่ด้วย (ดูตาราง "ที่มา" ใน docs/slide-script.md)
# =============================================================
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "docs" / "slide-assets"
OUT = ASSETS / "CopperDome-3min.pptx"

# Leelawadee UI มากับ Windows และมี glyph ไทยครบ — Tahoma เป็นตัวสำรอง
FONT = "Leelawadee UI"

NAVY = RGBColor(0x0F, 0x22, 0x33)
COPPER = RGBColor(0xB8, 0x73, 0x33)
INK = RGBColor(0x1C, 0x2B, 0x36)
MUTED = RGBColor(0x5A, 0x6B, 0x78)
BG = RGBColor(0xF7, 0xF8, 0xFA)
PANEL = RGBColor(0xE8, 0xEC, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.62)


def style(run, size, *, bold=False, color=INK, italic=False):
    """ตั้งฟอนต์ให้ครบทั้ง latin และ complex-script — ไม่ตั้ง cs แล้วไทยจะไม่ใช้ฟอนต์นี้"""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:cs", "a:ea"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return tf


def para(tf, text, size, *, bold=False, color=INK, space_after=8, first=False,
         align=PP_ALIGN.LEFT, italic=False, line=1.25, hang=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    if hang is not None:
        # แขวนย่อหน้า: บรรทัดที่ตัดคำต้องเยื้องมาอยู่ใต้ตัวอักษร ไม่ใช่ชิดขอบซ้าย
        ppr = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
        ppr.set("marL", str(int(hang)))
        ppr.set("indent", str(-int(hang)))
    style(p.add_run(), size, bold=bold, color=color, italic=italic)
    p.runs[0].text = text
    return p


def rect(slide, x, y, w, h, fill, *, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.5)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE and sh.adjustments:
        sh.adjustments[0] = radius
    sh.text_frame.word_wrap = True
    return sh


def new_slide(prs, title=None, *, tag=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    if title:
        tf = textbox(slide, MARGIN, Inches(0.42), W - 2 * MARGIN, Inches(0.75))
        para(tf, title, 32, bold=True, color=NAVY, first=True, space_after=0)
        rect(slide, MARGIN, Inches(1.24), Inches(1.5), Inches(0.075), COPPER,
             shape=MSO_SHAPE.RECTANGLE)
    if tag:
        tf = textbox(slide, W - MARGIN - Inches(3.6), Inches(0.55), Inches(3.6), Inches(0.4))
        para(tf, tag, 12, bold=True, color=COPPER, first=True, align=PP_ALIGN.RIGHT,
             space_after=0)
    return slide


def bullets(slide, x, y, w, items, size=17, gap=13):
    tf = textbox(slide, x, y, w, Inches(0.4))
    for i, item in enumerate(items):
        para(tf, "▸  " + item, size, color=INK, first=(i == 0), space_after=gap,
             hang=Inches(0.30))
    return tf


def picture(slide, name, x, y, width):
    """วางรูปโดยคุมความกว้าง แล้วคืนความสูงที่ได้ (คงสัดส่วน)"""
    path = ASSETS / name
    if not path.exists():
        raise FileNotFoundError(path)
    pic = slide.shapes.add_picture(str(path), x, y, width=width)
    return pic.height


def placeholder(slide, x, y, w, h, label, hint):
    rect(slide, x, y, w, h, PANEL, line=RGBColor(0xB6, 0xC0, 0xC8))
    tf = textbox(slide, x + Inches(0.3), y, w - Inches(0.6), h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "🖼  " + label, 16, bold=True, color=MUTED, first=True,
         align=PP_ALIGN.CENTER, space_after=6)
    para(tf, hint, 12, color=MUTED, align=PP_ALIGN.CENTER, space_after=0)


def callout(slide, x, y, w, h, lines, *, fill=COPPER, fg=WHITE):
    sh = rect(slide, x, y, w, h, fill)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.24)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (text, size, bold) in enumerate(lines):
        para(tf, text, size, bold=bold, color=fg, first=(i == 0), space_after=4)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- slides
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1 — ปก
    s = new_slide(prs)
    rect(s, Inches(0), Inches(0), Inches(0.28), H, COPPER, shape=MSO_SHAPE.RECTANGLE)
    tf = textbox(s, Inches(1.3), Inches(2.35), Inches(11), Inches(3), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "Copper Dome Turret", 56, bold=True, color=NAVY, first=True, space_after=14)
    para(tf, "ป้อมยิงอัตโนมัติ — คลิกเลือกเป้า แล้วป้อมหันไปยิงเอง", 24, color=INK,
         space_after=26)
    para(tf, "เขียนด้วย Python ทั้งหมด", 18, color=COPPER, bold=True, space_after=34)
    para(tf, "ARIS Project III  ·  Mini Project 1", 15, color=MUTED, space_after=4)
    para(tf, "[ใส่ชื่อสมาชิกทีม]", 15, color=MUTED, space_after=0)
    notes(s, "สวัสดีครับ กลุ่มเรานำเสนอ Copper Dome Turret — ป้อมยิงอัตโนมัติที่ให้ผู้ใช้"
             "คลิกเลือกเป้าบนจอ แล้วป้อมหันไปยิงเอง เขียนด้วย Python ทั้งหมด")

    # 2 — ภาพรวมระบบ
    s = new_slide(prs, "ภาพรวมระบบ")
    picture(s, "fig-block-diagram.png", MARGIN, Inches(1.95), W - 2 * MARGIN)
    callout(s, MARGIN, Inches(5.55), W - 2 * MARGIN, Inches(0.95),
            [("ตรรกะทุกบรรทัดอยู่ฝั่ง Python — Arduino เป็นแค่ช่องสัญญาณเข้า/ออก", 20, True)])
    notes(s, "ภาพรวมระบบเป็นแบบนี้ครับ กล้องมือถือส่งภาพเข้าคอม คอมรันโมเดลตรวจจับและ"
             "ตัวคุมการเล็ง แล้วส่งคำสั่งผ่าน Firmata ไปที่ Arduino ซึ่งทำหน้าที่เป็นแค่"
             "ช่องสัญญาณเข้าออก ไม่มีตรรกะอยู่บนบอร์ดเลย ตรรกะทุกบรรทัดอยู่ฝั่ง Python "
             "ตามกติกาที่กำหนดให้ใช้ Python อย่างเดียว")

    # 3 — กลไกการสร้าง
    s = new_slide(prs, "กลไกการสร้าง", tag="หัวข้อบังคับ ①")
    placeholder(s, MARGIN, Inches(1.62), Inches(5.5), Inches(4.05),
                "รูปถ่ายป้อมจริง", "ให้เห็น flywheel คู่ + แกน pan/tilt")
    bullets(s, Inches(6.5), Inches(1.72), Inches(6.2), [
        "กลไกยิง — ล้อ flywheel คู่หมุนสวนทาง หนีบลูกบอลไม้แล้วสะบัดออก",
        "ป้อนลูก — แมกกาซีน gravity-feed ไม่มีเซอร์โวดันลูก",
        "สองแกน — pan ±50°  ·  tilt ผ่านเฟืองสะพาน (rack) ช่วง 90–180°",
        "เซอร์โว MG945 สองตัว",
    ])
    callout(s, Inches(6.5), Inches(4.55), Inches(6.2), Inches(1.12), [
        ("⚙  ถอด gearbox ออกจากมอเตอร์ TT", 19, True),
        ("จากยิงไม่ออกเลย → ยิงได้เกิน 1.5 เมตร", 16, False),
    ])
    notes(s, "กลไกยิงเป็นล้อฟลายวีลสองล้อหมุนสวนทางกัน หนีบลูกบอลไม้แล้วสะบัดออก "
             "ป้อนลูกด้วยแมกกาซีนแรงโน้มถ่วง ไม่ต้องมีเซอร์โวดันลูก ตัวป้อมหมุนสองแกน "
             "แกนส่ายใช้เซอร์โว MG945 หมุนได้บวกลบห้าสิบองศา แกนเงยใช้เซอร์โวอีกตัว"
             "ขับผ่านเฟืองสะพาน — จุดที่ต้องเล่าคือมอเตอร์ครับ ตอนแรกใช้มอเตอร์ TT ทั้งชุด "
             "ยิงไม่ออกเลย เพราะมันมีเกียร์ทดรอบ เราถอดเกียร์บ็อกซ์ออกใช้แกนมอเตอร์เปลือย "
             "ยิงออกเกินหนึ่งเมตรครึ่งทันที เดี๋ยวสไลด์คำนวณจะบอกว่ารู้ได้ยังไงว่าต้องถอด")

    # 4 — โหมดสโคป (ไดอะแกรมวาดด้วย shape ไม่ใช่รูป)
    s = new_slide(prs, "แนวคิดออกแบบ: โหมดสโคป", tag="หัวข้อบังคับ ②")
    fx, fy, fw, fh = MARGIN, Inches(1.72), Inches(5.5), Inches(3.5)
    rect(s, fx, fy, fw, fh, WHITE, line=RGBColor(0x9A, 0xA7, 0xB1), shape=MSO_SHAPE.RECTANGLE)
    cx, cy = fx + fw / 2, fy + fh / 2
    arm = Inches(0.42)
    rect(s, cx - arm, cy - Emu(11430), arm * 2, Emu(22860), COPPER, shape=MSO_SHAPE.RECTANGLE)
    rect(s, cx - Emu(11430), cy - arm, Emu(22860), arm * 2, COPPER, shape=MSO_SHAPE.RECTANGLE)
    tgt = Inches(0.78)
    tx, ty = fx + Inches(3.75), fy + Inches(0.72)
    rect(s, tx, ty, tgt, tgt, RGBColor(0x2E, 0x86, 0xC1), shape=MSO_SHAPE.OVAL)
    rect(s, cx + Inches(0.32), cy - Inches(0.16), tx - cx - Inches(0.30), Inches(0.32),
         NAVY, shape=MSO_SHAPE.RIGHT_ARROW)
    rect(s, tx + tgt / 2 - Inches(0.16), ty + tgt + Inches(0.06), Inches(0.32),
         cy - ty - tgt - Inches(0.10), NAVY, shape=MSO_SHAPE.DOWN_ARROW)
    tf = textbox(s, cx + Inches(0.4), cy + Inches(0.22), Inches(1.7), Inches(0.3))
    para(tf, "error X", 13, bold=True, color=NAVY, first=True, space_after=0)
    tf = textbox(s, tx + tgt + Inches(0.12), ty + tgt + Inches(0.30), Inches(1.7), Inches(0.3))
    para(tf, "error Y", 13, bold=True, color=NAVY, first=True, space_after=0)
    tf = textbox(s, fx + Inches(0.2), fy + fh - Inches(0.52), fw - Inches(0.4), Inches(0.4))
    para(tf, "จุด zero (กากบาท) = ตำแหน่งที่ยิงแล้วโดน", 12, color=MUTED, first=True,
         space_after=0)
    tf = textbox(s, fx, fy + fh + Inches(0.16), fw, Inches(0.6))
    para(tf, "เล็ง = ลด error สองแกนให้เป็นศูนย์", 17, bold=True, color=COPPER, first=True,
         align=PP_ALIGN.CENTER, space_after=0)

    tf = textbox(s, Inches(6.5), Inches(1.72), Inches(6.2), Inches(1.1))
    para(tf, "กล้องติดไปกับลำกล้อง → ไม่ต้องวัดระยะถึงเป้าเลย", 19, bold=True,
         color=NAVY, first=True, space_after=0)
    bullets(s, Inches(6.5), Inches(2.72), Inches(6.2), [
        "P controller  ·  เกน 0.03 องศาต่อพิกเซล",
        "จำกัดก้าวละไม่เกิน 6 องศา กันเหวี่ยงเกิน",
        "ต้องอยู่ในกรอบ 15 พิกเซล ติดกัน 3 เฟรม ถึงประกาศว่าพร้อมยิง",
        "เป้าอยู่ระดับโต๊ะเดียวกัน + ยิงแรงคงที่ → จุด zero จุดเดียวคุมได้ทั้งช่วง 1.5–2.4 ม.",
    ])
    notes(s, "แนวคิดการเล็งครับ เราติดกล้องไว้กับลำกล้อง กล้องหันไปพร้อมป้อมเสมอ "
             "การเล็งจึงกลายเป็นการลดระยะห่างระหว่างเป้ากับจุดเล็งบนภาพให้เหลือศูนย์ "
             "ไม่ต้องวัดระยะถึงเป้าเลย ตัวคุมเป็น P controller ค่าเกน 0.03 องศาต่อพิกเซล "
             "จำกัดก้าวละไม่เกิน 6 องศากันเหวี่ยงเกิน และต้องอยู่ในกรอบ 15 พิกเซล"
             "ติดกันสามเฟรม ระบบถึงจะประกาศว่าพร้อมยิง")

    # 5 — วงจร + ทฤษฎี
    s = new_slide(prs, "วงจร และทฤษฎีที่ประยุกต์ใช้", tag="หัวข้อบังคับ ③")
    picture(s, "fig-circuit.png", MARGIN, Inches(1.95), Inches(7.15))
    bullets(s, Inches(8.05), Inches(1.72), Inches(4.68), [
        "18650 ×2 อนุกรม = 7.4 V แยกสองทาง",
        "→ L298N (H-bridge)  ·  PWM ที่ ENA/ENB คุมความเร็ว  ·  IN1–IN4 คุมทิศ",
        "→ buck 6 V เลี้ยงเซอร์โวแยก (ถ้าดึงจาก USB บอร์ดจะรีเซ็ตตัวเอง)",
        "กราวด์ทุกส่วนต่อร่วมกัน",
    ], size=15, gap=11)
    callout(s, Inches(8.05), Inches(4.72), Inches(4.68), Inches(1.68), [
        ("แบบจำลองกล้องรูเข็ม", 16, True),
        ("θ  =  arctan( Δpx / f )     ·     f = 1400 px", 15, True),
        ("แปลง “เป้าเยื้องกี่พิกเซล” → “ต้องหมุนกี่องศา”", 13, False),
    ], fill=NAVY)
    notes(s, "ฝั่งวงจรครับ แบตเตอรี่ 18650 สองก้อนอนุกรมได้ 7.4 โวลต์ แยกสองทาง "
             "ทางแรกเข้า L298N ซึ่งเป็นวงจร H-bridge คุมความเร็วมอเตอร์ด้วย PWM ที่ขา "
             "ENA ENB และคุมทิศด้วยขา IN หนึ่งถึงสี่ ทางที่สองผ่าน buck ลดเหลือ 6 โวลต์"
             "เลี้ยงเซอร์โวแยกต่างหาก เพราะถ้าดึงจากไฟ USB ของ Arduino บอร์ดจะรีเซ็ตตัวเอง "
             "กราวด์ทุกส่วนต่อร่วมกัน ส่วนทฤษฎีที่ใช้แปลงภาพเป็นมุมคือแบบจำลองกล้องรูเข็ม "
             "มุมเท่ากับอาร์กแทนเจนต์ของระยะพิกเซลหารความยาวโฟกัส ซึ่งเราวัดได้ 1400 พิกเซล")

    # 6 — ballistics
    s = new_slide(prs, "รายการคำนวณการเคลื่อนที่กระสุน", tag="หัวข้อบังคับ ④")
    eq = rect(s, MARGIN, Inches(1.62), W - 2 * MARGIN, Inches(1.02), WHITE,
              line=RGBColor(0xC8, 0xD0, 0xD6))
    eq.text_frame.margin_left = Inches(0.28)
    eq.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(eq.text_frame, "x(t) = v·cosθ·t          y(t) = h₀ + v·sinθ·t − ½gt²",
         20, bold=True, color=NAVY, first=True, space_after=4)
    para(eq.text_frame, "เป้าสูงกว่าปากกระบอก  Δh = +0.40 m   (โต๊ะ 800 − แท่นปืน 400 mm ตามโจทย์)",
         14, color=MUTED, space_after=0)

    rows = [
        ["ความเร็วต้นที่ต้องใช้ (m/s)", "1.50 m", "1.75 m", "2.00 m", "2.25 m", ""],
        ["θ = 25°", "6.70", "6.63", "6.70", "6.82", "⚠  ดิ่งแล้วขึ้น"],
        ["θ = 30°", "5.62", "5.73", "5.89", "6.07", "✓  เรียงขึ้น"],
        ["θ = 35°", "5.03", "5.21", "5.41", "5.61", "✓  เรียงขึ้น"],
    ]
    tw = Inches(7.55)
    tbl = s.shapes.add_table(4, 6, MARGIN, Inches(2.92), tw, Inches(1.85)).table
    for c, frac in enumerate([0.30, 0.125, 0.125, 0.125, 0.125, 0.20]):
        tbl.columns[c].width = Emu(int(tw * frac))
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(0.46)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (WHITE if r % 2 else PANEL)
            para(cell.text_frame, val, 13 if r == 0 else 14,
                 bold=(r == 0 or c == 0), color=(WHITE if r == 0 else INK), first=True,
                 align=PP_ALIGN.LEFT if c in (0, 5) else PP_ALIGN.CENTER, space_after=0)

    tf = textbox(s, Inches(8.5), Inches(2.92), Inches(4.25), Inches(3.2))
    hang = Inches(0.26)
    para(tf, "1.  ระยะโจทย์ 1.5–2.4 m → ต้องการความเร็วลูกแค่ ~5–6 m/s",
         14, first=True, space_after=11, hang=hang)
    para(tf, "2.  ที่ 25° ระยะเดียวได้จากความเร็ว 2 ค่า = กำกวมกลางช่วงใช้งาน "
             "→ ล็อกมุมเงย 30–35°", 14, space_after=11, hang=hang)
    para(tf, "3.  ความเร็วขอบล้อ ≈ 2× ความเร็วลูก · ล้อ Ø50 mm → ต้องหมุน ~3800 RPM",
         14, space_after=6, hang=hang)
    para(tf, "⟵  ตัวเลขนี้เองที่บอกว่ามอเตอร์มีเกียร์ (~200 RPM) ใช้ไม่ได้ ต้องถอด gearbox",
         14, bold=True, color=COPPER, space_after=0, hang=hang)

    tf = textbox(s, MARGIN, Inches(6.62), W - 2 * MARGIN, Inches(0.4))
    para(tf, "แบบจำลองไม่รวมแรงต้านอากาศ — ใช้กำหนดข้อกำหนดการออกแบบ "
             "ความแม่นจริงมาจากการยิงวัด", 11, color=MUTED, italic=True, first=True,
         space_after=0)
    notes(s, "รายการคำนวณการเคลื่อนที่กระสุนครับ เราใช้สมการโพรเจกไทล์มาตรฐาน "
             "แกนราบความเร็วคงที่ แกนดิ่งมีความเร่งโน้มถ่วง เป้าอยู่สูงกว่าปากกระบอก "
             "0.4 เมตรตามระยะสนามที่โจทย์กำหนด แก้สมการหาความเร็วต้นที่ต้องใช้ ได้ผลสามข้อ "
             "— หนึ่ง ระยะหนึ่งเมตรครึ่งถึงสองเมตรสี่ ต้องการความเร็วลูกแค่ราวห้าถึงหก"
             "เมตรต่อวินาที สอง ที่มุม 25 องศา ความเร็วกับระยะไม่เรียงกัน ระยะเดียวได้จาก"
             "ความเร็วสองค่า เราจึงเลือกล็อกมุมเงยที่ 30 ถึง 35 องศาแทน สาม ความเร็วขอบล้อ"
             "ประมาณสองเท่าของความเร็วลูก ล้อเส้นผ่านศูนย์กลาง 50 มิลลิเมตร ต้องหมุนราว"
             "สามพันแปดร้อยรอบต่อนาที — ตัวเลขนี้เองที่บอกเราว่ามอเตอร์ที่มีเกียร์ทดเหลือ"
             "สองร้อยรอบใช้ไม่ได้ ต้องถอดเกียร์ออก")

    # 7 — vision
    s = new_slide(prs, "ทำให้มองไม่มั่ว")
    picture(s, "fig-fp-rate.png", MARGIN, Inches(2.05), Inches(6.35))
    bullets(s, Inches(7.35), Inches(1.72), Inches(5.38), [
        "ชุดข้อมูลเก็บเอง 3,333 ภาพ  ·  3 ชนิด (capybara / dino / elephant)",
    ], size=16, gap=14)
    tf = textbox(s, Inches(7.35), Inches(2.62), Inches(5.38), Inches(0.4))
    para(tf, "ประตูกัน false positive 3 ชั้น", 17, bold=True, color=NAVY, first=True,
         space_after=0)
    bullets(s, Inches(7.35), Inches(3.12), Inches(5.38), [
        "① ความมั่นใจ ≥ 0.30",
        "② ขนาดกรอบต้องสมเหตุผลกับระยะ 0.5–4.5 ม.",
        "③ ต้องเห็นซ้ำที่เดิม 3 เฟรม",
    ], size=15, gap=10)
    callout(s, Inches(7.35), Inches(4.95), Inches(5.38), Inches(1.35), [
        ("บทเรียน", 14, True),
        ("ชุดวัดผลที่ไม่มีภาพพื้นหลังเลย = วัดการมองมั่วไม่ได้ตามนิยาม", 15, False),
    ], fill=NAVY)
    notes(s, "ฝั่งการมองเห็นครับ เราเก็บชุดข้อมูลเอง 3,333 ภาพ ปัญหาที่เจอคือโมเดลมอง"
             "ถุงพลาสติกเป็นคาปิบาร่าด้วยความมั่นใจ 0.83 สาเหตุคือชุดวัดผลไม่มีภาพ"
             "พื้นหลังเปล่าเลย แปลว่าเราวัดการมองมั่วไม่ได้ตั้งแต่ต้น เราเก็บภาพพื้นหลังเพิ่ม"
             "แล้วเทรนใหม่ อัตรามองมั่วบนห้องที่โมเดลไม่เคยเห็นลดจาก 8 เปอร์เซ็นต์เหลือ 1.1 "
             "โดยความไวไม่ตก\n\n"
             "*** ห้ามพูดว่าเป็นตัวเลขของห้องแข่ง — พูดได้แค่ 'ห้องที่โมเดลไม่เคยเห็น' ***")

    # 8 — UI
    s = new_slide(prs, "เลือกเป้าด้วย UI")
    placeholder(s, MARGIN, Inches(1.62), Inches(6.0), Inches(3.9),
                "ภาพหน้าจอขณะล็อกเป้า", "ให้เห็นกรอบตรวจจับ + เป้าเล็งแดง + จุด zero")
    bullets(s, Inches(7.0), Inches(1.75), Inches(5.72), [
        "คลิกซ้ายที่ตุ๊กตาตัวไหนก็ได้ใน 3 ตัว → ป้อมหันไปล็อกตัวนั้น",
        "F ยิง  ·  G หยุดล้อฉุกเฉิน (กดได้ทุกสถานะ)  ·  C คืนป้อมกลางลำ",
        "ขยับป้อมเองเมื่อไหร่ ระบบล้างสถานะล็อกทันที — กันยิงโดยเชื่อภาพเก่า",
    ], size=17, gap=15)
    callout(s, MARGIN, Inches(5.85), W - 2 * MARGIN, Inches(1.0),
            [("พร้อมสาธิตครับ", 28, True)])
    notes(s, "สุดท้ายคือส่วนติดต่อผู้ใช้ครับ ผู้ใช้คลิกที่ตุ๊กตาตัวไหนก็ได้ในสามตัว "
             "ป้อมจะหันไปล็อกตัวนั้น กด F ยิง และมีปุ่ม G หยุดล้อฉุกเฉินที่กดได้ทุกสถานะ "
             "พร้อมสาธิตแล้วครับ ขอบคุณครับ")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"saved: {OUT}")
    print(f"slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")


if __name__ == "__main__":
    build()
